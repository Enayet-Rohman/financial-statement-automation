from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

# === USER INPUT ===
ticker = input("Enter ticker (e.g., AAPL): ").upper()

# === LOAD TEMPLATE ===
wb = load_workbook("template.xlsx")
ws = wb["Data"]

ws["H1"] = f"TICKER: {ticker}"
ws["H2"] = "In $000's"

# === FIXED ROW MAP ===
rows = {
    "Revenue": 2,
    "Gross Profit": 3,
    "Operating Income": 4,
    "Net Income": 5,
    "EBITDA": 6,
    "Interest Expense": 7,
    "Total Assets": 8,
    "Total Equity": 9,
    "Current Assets": 10,
    "Current Liabilities": 11,
    "Total Debt": 12,
    "Cash & Equivalents": 13,
    "Operating Cash Flow": 14,
    "Capital Expenditure": 15
}

ciq_codes = {
    "Revenue": "IQ_TOTAL_REV",
    "Gross Profit": "IQ_GROSS_PROFIT",
    "Operating Income": "IQ_OPER_INC",
    "Net Income": "IQ_NET_INC",
    "EBITDA": "IQ_EBITDA",
    "Interest Expense": "IQ_INTEREST_EXP",
    "Total Assets": "IQ_TOTAL_ASSETS",
    "Total Equity": "IQ_TOTAL_EQUITY",
    "Current Assets": "IQ_TOTAL_CA",
    "Current Liabilities": "IQ_TOTAL_CL",
    "Total Debt": "IQ_TOTAL_DEBT",
    "Cash & Equivalents": "IQ_CASH_EQUIV",
    "Operating Cash Flow": "IQ_CASH_OPER",
    "Capital Expenditure": "IQ_CAPEX"
}

years = [2021, 2022, 2023, 2024, 2025]
year_cols = {2021:"B", 2022:"C", 2023:"D", 2024:"E", 2025:"F"}

# === FILL DATA SHEET WITH CIQ FORMULAS ===
for metric, row in rows.items():
    code = ciq_codes[metric]
    for y in years:
        col = year_cols[y]
        ws[f"{col}{row}"] = f'=CIQ("{ticker}","{code}","CY{y}")'

# === CREATE RATIOS SHEET ===
if "Ratios" in wb.sheetnames:
    rs = wb["Ratios"]
else:
    rs = wb.create_sheet("Ratios")

rs["A1"] = f"Financial Ratios for {ticker}"

# === RATIOS HEADER ===
rs["A2"] = "Ratio"
for i, y in enumerate(years):
    rs[f"{get_column_letter(2+i)}2"] = str(y)

ratio_row = 3

# === RATIO WRITER WITH FORMATTING ===
def write_ratio(name, formula_fn, is_percent=False, is_whole=False):
    global ratio_row
    rs[f"A{ratio_row}"] = name

    for i, y in enumerate(years):
        col = get_column_letter(2+i)
        cell = rs[f"{col}{ratio_row}"]
        cell.value = formula_fn(year_cols[y])

        # === APPLY FORMATTING ===
        if is_percent:
            cell.number_format = "0%"      # no decimals
        elif is_whole:
            cell.number_format = "0"       # whole number
        else:
            cell.number_format = "0.00"    # two decimals

    ratio_row += 1

# === MARGINS (percent, no decimals) ===
write_ratio("Gross Margin", lambda c: f"=Data!{c}{rows['Gross Profit']}/Data!{c}{rows['Revenue']}", True)
write_ratio("Operating Margin", lambda c: f"=Data!{c}{rows['Operating Income']}/Data!{c}{rows['Revenue']}", True)
write_ratio("Net Margin", lambda c: f"=Data!{c}{rows['Net Income']}/Data!{c}{rows['Revenue']}", True)
write_ratio("EBITDA Margin", lambda c: f"=Data!{c}{rows['EBITDA']}/Data!{c}{rows['Revenue']}", True)

# === LIQUIDITY (two decimals) ===
write_ratio("Current Ratio", lambda c: f"=Data!{c}{rows['Current Assets']}/Data!{c}{rows['Current Liabilities']}", False)
write_ratio("Cash Ratio", lambda c: f"=Data!{c}{rows['Cash & Equivalents']}/Data!{c}{rows['Current Liabilities']}", False)

# === LEVERAGE (two decimals) ===
write_ratio("Debt/Equity", lambda c: f"=Data!{c}{rows['Total Debt']}/Data!{c}{rows['Total Equity']}", False)
write_ratio("Debt/Assets", lambda c: f"=Data!{c}{rows['Total Debt']}/Data!{c}{rows['Total Assets']}", False)

# === FREE CASH FLOW (whole number) ===
write_ratio("Free Cash Flow", lambda c: f"=Data!{c}{rows['Operating Cash Flow']}+Data!{c}{rows['Capital Expenditure']}", False, True)

# === SAVE ===
wb.save("output.xlsx")
print("Done. Refresh CIQ in Excel.")
